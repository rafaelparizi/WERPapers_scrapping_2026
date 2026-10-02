"""Normalização de nomes de autores do dataset WERpapers.

Lê data/wer_papers.csv (gerado por scrape_werpapers.py) e grava:
  data/authors_mapping.csv    variante -> nome canônico (para revisão manual)
  data/wer_authorships.csv    uma linha por (artigo, autor), formato longo para bibliometria
  data/wer_papers.{csv,xlsx,json} com a coluna nova `authors_normalized`

Correções manuais: data/author_overrides.csv (colunas variant,canonical,paper_id) é aplicado por
último e prevalece sobre o agrupamento automático. Use canonical vazio para impedir que uma variante
seja agrupada (ela fica como está). Preencha paper_id para aplicar a correção só naquele artigo
(homônimos com a mesma grafia).
"""
import json
import re
import unicodedata
from collections import Counter
from pathlib import Path

import pandas as pd

DATA = Path(__file__).parent / "data"
PARTICLES = {"de", "da", "do", "dos", "das", "del", "della", "la", "las", "los", "le", "e", "y", "van", "von", "der", "di", "du"}
SUFFIXES = {"jr", "junior", "neto", "filho", "sobrinho"}
# Primeiros nomes que costumam ser omitidos ("María Carmen Leonardi" -> "Carmen Leonardi")
SKIPPABLE_FIRST = {"maria", "jose", "juan", "joao", "ana", "luis", "luiz", "carlos", "jorge"}


# ---------------------------------------------------------------- limpeza

def fix_case(token: str) -> str:
    """'LEONARDI' -> 'Leonardi', 'CastrO' -> 'Castro', 'LErário' -> 'Lerário'; preserva iniciais ('F.A.')
    e trata cada parte de nomes com hífen ('Bulcão-Neto')."""
    if "-" in token.strip("-"):
        return "-".join(fix_case(t) for t in token.split("-"))
    letters = re.sub(r"[^A-Za-zÀ-ÿ]", "", token)
    if len(letters) <= 1 or re.fullmatch(r"([A-ZÀ-Þ]\.)+-?", token):
        return token
    if letters[1:] != letters[1:].lower():
        return token[0].upper() + token[1:].lower()
    return token


def split_authors(raw: str) -> list[str]:
    raw = raw.replace("\xa0", " ").replace("´", "")
    raw = re.sub(r"\s+(?:and|e)\s+(?=[A-ZÀ-Ý])", ";", raw)  # "Gomes e Patrícia" / "Fabri and Alexandre"
    raw = re.sub(r";\s*e\s+", ";", raw)                       # "; e Marcia Lucena"
    names = []
    for part in raw.split(";"):
        name = re.sub(r"[\d*]+$", "", part.strip()).strip().rstrip(".").strip()
        name = re.sub(r"\s+", " ", name)
        if not name:
            continue
        if re.fullmatch(r"(?i)jr|júnior|junior", name) and names:  # "João Batista Camargo; Jr."
            names[-1] += " Jr."
            continue
        name = " ".join(fix_case(t) for t in name.split())
        names.append(name)
    return names


# ---------------------------------------------------------------- comparação

def fold(s: str) -> str:
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()


def tokens(name: str) -> list[str]:
    """Tokens dobrados, sem partículas nem sufixos; iniciais viram tokens de 1 letra ('F.A.' -> 'f','a')."""
    out = []
    for t in re.split(r"[\s\-]+", fold(name)):
        parts = [p for p in t.split(".") if p] if "." in t else [t]
        for p in parts:
            p = re.sub(r"[^a-z]", "", p)
            if p and p not in PARTICLES and p not in SUFFIXES:
                out.append(p)
    return out


def tok_match(a: str, b: str) -> bool:
    if len(a) == 1 or len(b) == 1:
        return a[0] == b[0]
    return a == b


def embeds(short: list[str], long: list[str]) -> bool:
    """Cada token de `short` casa, em ordem, com um token de `long`; o último token de `short`
    (sobrenome) precisa estar por extenso e casar por extenso."""
    if not short or len(short[-1]) == 1:
        return False
    i = 0
    for k, s in enumerate(short):
        while i < len(long) and not tok_match(s, long[i]):
            i += 1
        if i == len(long):
            return False
        if k == len(short) - 1 and long[i] != s:
            return False
        i += 1
    return True


def compatible(a: list[str], b: list[str]) -> str | None:
    """Retorna a regra que justifica a união, ou None."""
    short, long = sorted((a, b), key=lambda t: (len(t), sum(map(len, t))))
    if short == long:
        return "same_key"
    if len(short) < 2:  # nome de um só token ("Rosângela De") não é unido automaticamente
        return None
    if not tok_match(short[0], long[0]):
        # "Carmen Leonardi" ~ "María Carmen Leonardi"
        if long[0] in SKIPPABLE_FIRST and len(long) > 2 and short[0] == long[1] and embeds(short, long[1:]):
            return "first_name_omitted"
        return None
    if not embeds(short, long):
        return None
    if short[-1] != long[-1]:  # "Oscar Pastor" ~ "Oscar Pastor López", mas também "Ricardo Almeida" ~ "Ricardo de Almeida Falbo"
        return "surname_not_last"
    if len(short[0]) == 1 or len(long[0]) == 1:
        return "initials"
    return "subset"


# ---------------------------------------------------------------- agrupamento

def build_clusters(counts: Counter) -> dict[str, dict]:
    # 1) junta grafias com a mesma chave dobrada (acentos/caixa/pontuação)
    by_key: dict[tuple, list[str]] = {}
    for name in counts:
        by_key.setdefault(tuple(tokens(name)), []).append(name)
    key_count = {k: sum(counts[n] for n in ns) for k, ns in by_key.items()}

    # 2) do nome mais completo ao menos completo; cada chave entra no grupo cujo fundador
    #    (a forma mais completa) é compatível com ela. Compatível com 2+ fundadores = ambígua.
    keys = sorted(by_key, key=lambda k: (-sum(len(x) > 1 for x in k), -len(k), -key_count[k]))
    founders: list[tuple] = []
    members: dict[tuple, list[tuple]] = {}
    rule_of: dict[tuple, str] = {}
    for k in keys:
        hits = [(f, r) for f in founders if (r := compatible(k, f))]
        if len(hits) == 1:
            f, rule_of[k] = hits[0]
            members[f].append(k)
        else:
            founders.append(k)
            members[k] = [k]
            rule_of[k] = "ambiguous" if hits else ""

    result = {}
    for f, ks in members.items():
        # canônico: chave mais frequente entre as que têm primeiro nome por extenso;
        # dentro dela, a grafia mais frequente (empate -> com acento)
        full = [k for k in ks if len(k[0]) > 1] or ks
        best = max(full, key=lambda k: (key_count[k], len(k)))
        canonical = max(by_key[best], key=lambda n: (counts[n], n != fold(n)))
        n_var = sum(len(by_key[k]) for k in ks)
        for k in ks:
            for n in by_key[k]:
                rule = rule_of[k] or ("same_key" if k == f and n != canonical else "")
                if k == f and k != best and rule != "ambiguous":
                    rule = compatible(k, best) or "subset"
                if n == canonical and rule != "ambiguous":
                    rule = ""
                result[n] = {"canonical": canonical, "rule": rule, "cluster_variants": n_var}
    return result


# ---------------------------------------------------------------- planilha com destaque

def write_xlsx(out: pd.DataFrame, names_col: pd.Series, mapping: dict, review: set, mp: pd.DataFrame):
    from openpyxl import Workbook
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles import Alignment, Font, PatternFill

    fills = {"revisar": PatternFill("solid", fgColor="F8D7DA"), "ajustado": PatternFill("solid", fgColor="FFF3CD")}
    font_review = InlineFont(b=True, color="C00000")
    font_adjusted = InlineFont(b=True, color="9C5700")

    def highlighted(raw: str, names: list[str], status: str):
        """Texto original de `authors` com os nomes afetados em negrito/cor."""
        spans = []
        pos = 0
        for n in names:
            kind = "review" if n in review and status == "revisar" else "adjusted" if n != mapping[n]["canonical"] else None
            if kind is None:
                continue
            m = re.search(re.escape(re.sub(r" Jr\.$", "", n)), raw[pos:], re.I)
            if m:
                spans.append((pos + m.start(), pos + m.end(), kind))
                pos += m.end()
        if not spans:
            return raw
        parts, last = [], 0
        for a, b, kind in spans:
            if a > last:
                parts.append(raw[last:a])
            parts.append(TextBlock(font_review if kind == "review" else font_adjusted, raw[a:b]))
            last = b
        if last < len(raw):
            parts.append(raw[last:])
        return CellRichText(parts)

    wb = Workbook()
    ws = wb.active
    ws.title = "papers"
    cols = list(out.columns)
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    ai, si = cols.index("authors"), cols.index("name_status")
    for (_, row), names in zip(out.iterrows(), names_col):
        vals = [v.item() if hasattr(v, "item") else v for v in row.tolist()]
        ws.append(vals)
        r = ws.max_row
        status = vals[si]
        ws.cell(r, ai + 1).value = highlighted(vals[ai], names, vals[si])
        if status in fills:
            for c in ws[r]:
                c.fill = fills[status]
    widths = {"title": 60, "authors": 60, "authors_normalized": 60, "name_changes": 60, "abstract": 50,
              "keywords": 30, "track": 25, "edition_name": 30, "location_date": 25, "paper_url": 25, "pdf_url": 25}
    for idx, col in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = widths.get(col, 12)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    rv = wb.create_sheet("revisar_nomes")
    rv.append(["variant", "canonical", "rule", "shared_coauthors", "papers"])
    papers_of: dict[str, list] = {}
    for pid, names in zip(out["paper_id"], names_col):
        for n in names:
            papers_of.setdefault(n, []).append(pid)
    for _, m in mp[mp.needs_review].iterrows():
        rv.append([m.variant, m.canonical, m.rule, m.shared_coauthors, ", ".join(papers_of.get(m.variant, []))])
    for c in rv[1]:
        c.font = Font(bold=True)
    for letter, w in zip("ABCDE", (35, 40, 18, 16, 30)):
        rv.column_dimensions[letter].width = w

    lg = wb.create_sheet("legenda")
    for line in [
        ("Situação dos nomes (coluna name_status)",),
        ("revisar", "Linha em vermelho claro. Há nome cuja normalização é incerta: sem coautor em comum com as demais variantes, união por iniciais, sobrenome que não é o último, ou variante ambígua. O nome aparece em vermelho e negrito na coluna authors."),
        ("ajustado", "Linha em amarelo. Algum nome foi normalizado ou a separação dos autores foi corrigida. O nome aparece em laranja e negrito na coluna authors."),
        ("ok", "Sem alteração."),
        (),
        ("name_changes", "Lista 'original → canônico [regra]' dos nomes alterados no artigo."),
        ("revisar_nomes", "Aba com cada nome a revisar e os artigos em que aparece."),
        ("Como corrigir", "Adicionar linha em data/author_overrides.csv (variant,canonical) e rodar normalize_authors.py. canonical vazio impede a união."),
    ]:
        lg.append(line)
    lg["A1"].font = Font(bold=True)
    lg["A2"].fill, lg["A3"].fill = fills["revisar"], fills["ajustado"]
    lg.column_dimensions["A"].width = 18
    lg.column_dimensions["B"].width = 120
    for row in lg.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(DATA / "wer_papers.xlsx")


def write_normalized_xlsx(out: pd.DataFrame, authorships: pd.DataFrame):
    """Planilha final para análise: sem prefácios, coluna authors = nomes normalizados."""
    from openpyxl.styles import Font

    papers = out[out["is_preface"].astype(str) != "True"].copy()
    papers = papers.rename(columns={"authors": "authors_original"}).rename(columns={"authors_normalized": "authors"})
    cols = ["paper_id", "year", "edition", "edition_name", "location_date", "track", "title", "authors",
            "n_authors", "abstract", "keywords", "doi", "pdf_url", "paper_url", "issn", "isbn", "publisher",
            "authors_original"]
    widths = {"title": 60, "authors": 60, "authors_original": 60, "abstract": 50, "keywords": 30, "track": 25,
              "edition_name": 30, "location_date": 25, "paper_url": 25, "pdf_url": 25, "author": 35}
    with pd.ExcelWriter(DATA / "wer_papers_normalized.xlsx", engine="openpyxl") as xw:
        papers[cols].to_excel(xw, sheet_name="papers", index=False)
        authorships[["paper_id", "year", "position", "author", "author_original", "title"]].to_excel(
            xw, sheet_name="authorships", index=False)
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True)
                ws.column_dimensions[c.column_letter].width = widths.get(c.value, 12)
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions
    return len(papers)


def main():
    df = pd.read_csv(DATA / "wer_papers.csv", keep_default_na=False)
    df["is_preface"] = df["is_preface"].astype(str) == "True"
    df["names_"] = df["authors"].map(split_authors)

    counts = Counter(n for names in df.loc[~df.is_preface, "names_"] for n in names)
    for names in df.loc[df.is_preface, "names_"]:
        counts.update(n for n in names if n not in counts)
    mapping = build_clusters(counts)

    overrides_path = DATA / "author_overrides.csv"
    if not overrides_path.exists():
        overrides_path.write_text("variant,canonical,paper_id\n", encoding="utf-8")
    paper_ov: dict[tuple[str, str], str] = {}  # (paper_id, variant) -> canonical
    for _, o in pd.read_csv(overrides_path, keep_default_na=False, dtype=str).iterrows():
        if o.variant not in mapping:
            print(f"Aviso: variante do override não encontrada: {o.variant!r}")
        elif o.get("paper_id", ""):
            paper_ov[(o.paper_id, o.variant)] = o.canonical or o.variant
        else:
            mapping[o.variant].update(canonical=o.canonical or o.variant, rule="manual")

    def canon(n: str, pid: str | None = None) -> str:
        return paper_ov.get((pid, n), mapping[n]["canonical"])

    df["authors_normalized"] = [
        "; ".join(dict.fromkeys(canon(n, pid) for n in ns)) for pid, ns in zip(df["paper_id"], df["names_"])]
    df["n_authors"] = df["authors_normalized"].map(lambda s: len([a for a in s.split("; ") if a]))

    # formato longo
    rows = [{"paper_id": r.paper_id, "year": r.year, "position": i + 1, "author_original": n,
             "author": canon(n, r.paper_id), "title": r.title}
            for r in df[~df.is_preface].itertuples() for i, n in enumerate(r.names_)]
    authorships = pd.DataFrame(rows)
    authorships.to_csv(DATA / "wer_authorships.csv", index=False, encoding="utf-8-sig")

    # tabela de mapeamento para revisão
    papers_per = Counter(r["author_original"] for r in rows)
    mp = pd.DataFrame([{"variant": n, "canonical": m["canonical"], "rule": m["rule"],
                        "n_papers_variant": papers_per.get(n, 0), "cluster_variants": m["cluster_variants"]}
                       for n, m in mapping.items()])
    # evidência por coautoria: a variante tem coautor em comum com as outras variantes do grupo?
    coauth: dict[str, set] = {}
    for pid, names in zip(df.loc[~df.is_preface, "paper_id"], df.loc[~df.is_preface, "names_"]):
        for n in names:
            coauth.setdefault(n, set()).update(canon(o, pid) for o in names if o != n)
    by_canon = mp.groupby("canonical")["variant"].apply(list).to_dict()
    mp["shared_coauthors"] = [
        len(coauth.get(v, set()) & set().union(*(coauth.get(o, set()) for o in by_canon[c] if o != v)))
        if v != c else "" for v, c in zip(mp.variant, mp.canonical)]
    risky = {"initials", "first_name_omitted", "surname_not_last", "ambiguous", "subset"}
    mp["needs_review"] = [r in risky and (r == "ambiguous" or s == 0)
                          for r, s in zip(mp.rule, mp.shared_coauthors)]
    mp = mp.sort_values(["canonical", "n_papers_variant"], ascending=[True, False])
    mp.to_csv(DATA / "authors_mapping.csv", index=False, encoding="utf-8-sig")

    # situação dos nomes em cada artigo
    review = set(mp.loc[mp.needs_review, "variant"])
    statuses, changes = [], []
    for pid, raw, names in zip(df["paper_id"], df["authors"], df["names_"]):
        items = []
        for n in names:
            m = mapping[n]
            if (pid, n) in paper_ov:
                items.append(f"{n} → {paper_ov[(pid, n)]} [manual]")
            elif n in review:
                items.append(f"{n} → {m['canonical']} [revisar: {m['rule']}]" if n != m["canonical"]
                             else f"{n} [revisar: {m['rule']}]")
            elif n != m["canonical"]:
                items.append(f"{n} → {m['canonical']} [{m['rule']}]")
        if [x.strip() for x in raw.split(";") if x.strip()] != names:
            items.insert(0, "[separação/grafia dos autores corrigida]")
        changes.append("; ".join(items))
        statuses.append("revisar" if any(n in review and (pid, n) not in paper_ov for n in names)
                        else "ajustado" if items else "ok")
    df["name_status"] = statuses
    df["name_changes"] = changes

    cols = [c for c in df.columns if c != "names_"]
    i = cols.index("authors")
    for k, c in enumerate(["authors_normalized", "name_status", "name_changes"], start=1):
        cols.insert(i + k, cols.pop(cols.index(c)))
    out = df[cols]
    out.to_csv(DATA / "wer_papers.csv", index=False, encoding="utf-8-sig")
    write_xlsx(out.drop(columns="bibtex"), df["names_"], mapping, review, mp)
    n_final = write_normalized_xlsx(out, authorships)
    (DATA / "wer_papers.json").write_text(json.dumps(out.to_dict("records"), ensure_ascii=False, indent=2), encoding="utf-8")

    merged = mp[mp.variant != mp.canonical]
    print(f"Grafias originais: {len(mapping)} | Autores após normalização: {authorships.author.nunique()}")
    print(f"Variantes unidas: {len(merged)} | por regra: {merged.rule.value_counts().to_dict()}")
    print(f"Ambíguas (não unidas): {(mp.rule == 'ambiguous').sum()} | Marcadas para revisão: {mp.needs_review.sum()}")
    print(f"Artigos por situação dos nomes: {out.name_status.value_counts().to_dict()}")
    print(f"wer_papers_normalized.xlsx: {n_final} artigos (sem prefácios)")


if __name__ == "__main__":
    main()
